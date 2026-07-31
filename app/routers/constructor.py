"""Konstruktor — boshqaruvchi o'ziga kerakli qo'shimcha bo'limlarni o'zi yaratadi.

Masalan: "Transport xarajatlari", "Stanoklar ro'yxati", "Qarz daftari" —
nom + maydonlar (matn/raqam/pul/sana) belgilanadi, keyin yozuvlar kiritiladi
va 1 klikda Excel qilib olinadi.
"""
import io
import json
from datetime import date
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pydantic import BaseModel
from sqlalchemy.orm import Session
from ..db import get_db
from ..auth import get_user, require_roles
from .. import models as m

router = APIRouter(prefix="/api/sections", tags=["Konstruktor"])

FIELD_TYPES = ["matn", "raqam", "pul", "sana"]


class FieldIn(BaseModel):
    label: str
    type: str = "matn"


class SectionIn(BaseModel):
    name: str
    icon: str = "📋"
    fields: list[FieldIn]


class RecordIn(BaseModel):
    data: dict


def section_out(db: Session, s: m.CustomSection) -> dict:
    count = db.query(m.CustomRecord).filter(m.CustomRecord.section_id == s.id).count()
    return {"id": s.id, "name": s.name, "icon": s.icon,
            "fields": json.loads(s.fields_json), "records_count": count}


@router.get("")
def list_sections(db: Session = Depends(get_db), user=Depends(get_user)):
    return [section_out(db, s) for s in db.query(m.CustomSection).order_by(m.CustomSection.id).all()]


@router.post("")
def create_section(data: SectionIn, db: Session = Depends(get_db),
                   user=Depends(require_roles("Rahbar"))):
    if not data.name.strip():
        raise HTTPException(400, "Bo'lim nomi bo'sh bo'lmasin")
    if not data.fields:
        raise HTTPException(400, "Kamida 1 ta maydon kerak")
    fields = []
    for i, f in enumerate(data.fields):
        if f.type not in FIELD_TYPES:
            raise HTTPException(400, f"Maydon turi noto'g'ri: {f.type}. Ruxsat: {FIELD_TYPES}")
        if not f.label.strip():
            raise HTTPException(400, "Maydon nomi bo'sh bo'lmasin")
        fields.append({"key": f"f{i+1}", "label": f.label.strip(), "type": f.type})
    s = m.CustomSection(name=data.name.strip(), icon=data.icon or "📋",
                        fields_json=json.dumps(fields, ensure_ascii=False))
    db.add(s)
    db.add(m.AuditLog(who=user.name, action="Yangi bo'lim yaratildi (konstruktor)",
                      detail=data.name))
    db.commit()
    return section_out(db, s)


@router.delete("/{sid}")
def delete_section(sid: int, db: Session = Depends(get_db),
                   user=Depends(require_roles("Rahbar"))):
    s = db.get(m.CustomSection, sid)
    if not s:
        raise HTTPException(404, "Bo'lim topilmadi")
    db.query(m.CustomRecord).filter(m.CustomRecord.section_id == sid).delete()
    db.delete(s)
    db.add(m.AuditLog(who=user.name, action="Bo'lim o'chirildi", detail=s.name))
    db.commit()
    return {"ok": True}


@router.get("/{sid}/records")
def list_records(sid: int, db: Session = Depends(get_db), user=Depends(get_user)):
    s = db.get(m.CustomSection, sid)
    if not s:
        raise HTTPException(404, "Bo'lim topilmadi")
    rows = (db.query(m.CustomRecord).filter(m.CustomRecord.section_id == sid)
            .order_by(m.CustomRecord.id.desc()).limit(500).all())
    return {"section": section_out(db, s),
            "records": [{"id": r.id, "data": json.loads(r.data_json),
                         "created_at": r.created_at.isoformat()} for r in rows]}


@router.post("/{sid}/records")
def add_record(sid: int, data: RecordIn, db: Session = Depends(get_db),
               user=Depends(get_user)):
    s = db.get(m.CustomSection, sid)
    if not s:
        raise HTTPException(404, "Bo'lim topilmadi")
    fields = json.loads(s.fields_json)
    clean = {}
    for f in fields:
        v = data.data.get(f["key"], "")
        if f["type"] in ("raqam", "pul") and v not in ("", None):
            try:
                v = float(v)
            except (TypeError, ValueError):
                raise HTTPException(400, f"«{f['label']}» raqam bo'lishi kerak")
        clean[f["key"]] = v
    r = m.CustomRecord(section_id=sid, data_json=json.dumps(clean, ensure_ascii=False))
    db.add(r)
    db.commit()
    return {"id": r.id}


@router.delete("/{sid}/records/{rid}")
def delete_record(sid: int, rid: int, db: Session = Depends(get_db),
                  user=Depends(require_roles("Rahbar"))):
    r = db.get(m.CustomRecord, rid)
    if not r or r.section_id != sid:
        raise HTTPException(404, "Yozuv topilmadi")
    db.delete(r)
    db.commit()
    return {"ok": True}


@router.get("/{sid}/export.xlsx")
def export_section(sid: int, db: Session = Depends(get_db), user=Depends(get_user)):
    s = db.get(m.CustomSection, sid)
    if not s:
        raise HTTPException(404, "Bo'lim topilmadi")
    fields = json.loads(s.fields_json)
    wb = Workbook()
    ws = wb.active
    ws.title = s.name[:30]
    ws.append(["№", "Sana"] + [f["label"] for f in fields])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5E63E0")
        cell.alignment = Alignment(horizontal="center")
    rows = (db.query(m.CustomRecord).filter(m.CustomRecord.section_id == sid)
            .order_by(m.CustomRecord.id).all())
    for i, r in enumerate(rows, 1):
        d = json.loads(r.data_json)
        ws.append([i, r.created_at.strftime("%d.%m.%Y")] + [d.get(f["key"], "") for f in fields])
    for i, f in enumerate(fields):
        col = chr(ord("C") + i)
        ws.column_dimensions[col].width = 20
        if f["type"] == "pul":
            for cc in ws[col]:
                cc.number_format = "#,##0"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="bolim_{sid}_{date.today()}.xlsx"'})
