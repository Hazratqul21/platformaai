"""Kassa jurnali — kirim-chiqim daftari (приход-расход Excel'ning o'rni).

Har bir pul harakati bitta jurnalda: kimdan keldi / kimga-nimaga ketdi.
Sex/filial bo'yicha ajratiladi (Excel'da 2 sheet = 2 firma edi).
"""
import io
from datetime import date, datetime
from decimal import Decimal
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pydantic import BaseModel
from sqlalchemy import func
from sqlalchemy.orm import Session
from ..db import get_db
from ..auth import get_user, require_roles
from .. import models as m

router = APIRouter(prefix="/api/kassa", tags=["Kassa jurnali"])
kassachi = require_roles("Rahbar", "Buxgalter")


class EntryIn(BaseModel):
    direction: str            # Kirim / Chiqim
    who: str
    note: str = ""
    amount: float
    currency: str = "so'm"
    firm: str = "Asosiy"
    entry_at: str | None = None   # YYYY-MM-DD (bo'sh = bugun)


def entry_out(e: m.KassaEntry) -> dict:
    return {"id": e.id, "firm": e.firm, "direction": e.direction, "who": e.who,
            "note": e.note, "amount": float(e.amount), "currency": e.currency,
            "entry_at": e.entry_at.isoformat(), "created_by": e.created_by}


@router.get("/firms")
def firms(db: Session = Depends(get_db), user=Depends(get_user)):
    """Barcha firmalar ro'yxati (kassa + mijoz + xodim + xariddan yig'iladi)."""
    names = set()
    for model, col in ((m.KassaEntry, m.KassaEntry.firm), (m.Client, m.Client.firm),
                       (m.Employee, m.Employee.firm), (m.Purchase, m.Purchase.firm)):
        for (v,) in db.query(col).distinct().all():
            if v:
                names.add(v)
    return sorted(names)


@router.get("")
def journal(firm: str | None = None, db: Session = Depends(get_db), user=Depends(get_user)):
    # Firmasi ko'rsatilmagan yozuv ikkala sex ostida ham ko'rinadi — mijoz,
    # buyurtma va xarajat ro'yxatlaridagi qoida bilan bir xil. Aks holda avtomat
    # yozuvlar (firmasi belgilanmagan mijozning to'lovi) filtrda yo'qolib ketardi.
    def fk(q):
        return q.filter(m.KassaEntry.firm.in_([firm, ""])) if firm else q

    rows = (fk(db.query(m.KassaEntry))
            .order_by(m.KassaEntry.entry_at.desc(), m.KassaEntry.id.desc()).limit(600).all())

    firms = [r[0] for r in db.query(m.KassaEntry.firm).distinct().all() if r[0]]

    def sums(direction, cur="so'm"):
        qq = fk(db.query(func.coalesce(func.sum(m.KassaEntry.amount), 0)).filter(
            m.KassaEntry.direction == direction, m.KassaEntry.currency == cur))
        return float(qq.scalar())

    kirim, chiqim = sums("Kirim"), sums("Chiqim")
    usd_k, usd_c = sums("Kirim", "USD"), sums("Chiqim", "USD")
    return {
        "entries": [entry_out(e) for e in rows],
        "firms": firms,
        "kirim": kirim, "chiqim": chiqim, "balans": kirim - chiqim,
        "usd_kirim": usd_k, "usd_chiqim": usd_c,
    }


@router.post("")
def add_entry(d: EntryIn, db: Session = Depends(get_db), user=Depends(kassachi)):
    if d.direction not in ("Kirim", "Chiqim"):
        raise HTTPException(400, "Yo'nalish: Kirim yoki Chiqim")
    if not d.who.strip():
        raise HTTPException(400, "Kimga/kimdan maydoni bo'sh bo'lmasin")
    if d.amount <= 0:
        raise HTTPException(400, "Сумма 0 дан катта бўлсин")
    if d.currency not in ("so'm", "USD"):
        raise HTTPException(400, "Valyuta: so'm yoki USD")
    when = date.today()
    if d.entry_at:
        try:
            when = datetime.strptime(d.entry_at, "%Y-%m-%d").date()
        except ValueError:
            raise HTTPException(400, "Sana formati: YYYY-MM-DD")
    e = m.KassaEntry(firm=d.firm.strip() or "Asosiy", direction=d.direction,
                     who=d.who.strip(), note=d.note.strip(),
                     amount=Decimal(str(d.amount)), currency=d.currency,
                     entry_at=when, created_by=user.name)
    db.add(e)
    db.add(m.AuditLog(who=user.name, action=f"Kassa: {d.direction.lower()}",
                      detail=f"{d.who} · {d.amount:,.0f} {d.currency}".replace(",", " ")))
    db.commit()

    out = entry_out(e)
    # Chiqim "kim"i yetkazib beruvchi nomiga o'xshasa — ikki marta yozib
    # qo'ymaslik uchun ogohlantiramiz (kassa va Xarid alohida hisoblar)
    if d.direction == "Chiqim":
        # SQLite lower() kirillni tushunmaydi — Python'da solishtiramiz
        nom = d.who.strip().lower()
        sup = None
        if nom:
            for su in db.query(m.Supplier).all():
                sn = su.name.lower()
                if nom in sn or sn in nom:
                    sup = su
                    break
        if sup:
            out["ogoh"] = (f"«{sup.name}» — бу етказиб берувчи. Агар унга берган пулни "
                           f"Харид бўлимида ҳам ёзган бўлсангиз, бу пул икки марта "
                           f"ҳисобланмаслиги учун бу ерга такрор ёзманг. Етказиб берувчига "
                           f"тўловни Омбор > Пул бериш орқали юритинг.")
    return out


@router.delete("/{eid}")
def del_entry(eid: int, db: Session = Depends(get_db), user=Depends(require_roles("Rahbar"))):
    e = db.get(m.KassaEntry, eid)
    if not e:
        raise HTTPException(404, "Yozuv topilmadi")
    # Bog'langan manba o'z bo'limidan o'chirilishi kerak — aks holda mijoz qarzı
    # yoki xarid qarzi kassadan mustaqil o'zgarib ketadi.
    if e.linked_payment_id:
        raise HTTPException(400, "Бу ёзув мижоз тўловига боғланган — Молия > Тўловлардан ўчиринг")
    if e.linked_purchase_payment_id or e.linked_purchase_id:
        raise HTTPException(400, "Бу ёзув харид тўловига боғланган — Харид бўлимидан ўчиринг")
    if e.linked_supplier_payment_id:
        raise HTTPException(400, "Бу ёзув етказиб берувчи тўловига боғланган — Омбордан ўчиринг")
    # sex xarajatidan / avansdan avtomat yaratilgan bo'lsa — bog'langan cash_entry
    # ham o'chadi (takror qolmasin: Цех харажатлари / Avans ro'yxatidan ham ketadi)
    if e.linked_cash_id:
        ce = db.get(m.CashEntry, e.linked_cash_id)
        if ce:
            db.delete(ce)
    db.delete(e)
    db.add(m.AuditLog(who=user.name, action="Kassa yozuvi o'chirildi",
                      detail=f"{e.direction} · {e.who} · {float(e.amount):,.0f}".replace(",", " ")))
    db.commit()
    return {"ok": True}


@router.get("/export.xlsx")
def export_xlsx(firm: str | None = None, db: Session = Depends(get_db), user=Depends(get_user)):
    """Kirim-chiqim Excel — приход-расход fayl formatida (chapda chiqim, o'ngda kirim)."""
    HDR = Font(bold=True, color="FFFFFF")
    FILL = PatternFill("solid", fgColor="5E63E0")
    wb = Workbook()
    # Bitta firma so'ralganda — ekrandagi ro'yxat bilan bir xil bo'lsin: firmasi
    # belgilanmagan yozuvlar ham qo'shiladi. Hamma firma so'ralganda esa har yozuv
    # aynan bitta varaqqa tushadi (aks holda jamilar ikki marta sanalardi), firmasi
    # belgilanmaganlari «Умумий» varag'ida yig'iladi.
    if firm:
        varaqlar = [(firm, m.KassaEntry.firm.in_([firm, ""]))]
    else:
        nomlar = sorted({r[0] or "" for r in db.query(m.KassaEntry.firm).distinct().all()},
                        key=lambda x: (x == "", x)) or ["Asosiy"]
        varaqlar = [(nom, m.KassaEntry.firm == nom) for nom in nomlar]
    first = True
    for fm, shart in varaqlar:
        ws = wb.active if first else wb.create_sheet()
        ws.title = (fm[:30] or "Умумий")
        first = False
        ws.append(["Sana", "Kimga - Nimaga", "Izoh", "Xarajat (chiqim)",
                   "Kimdan", "Izoh", "Kirim"])
        for cell in ws[1]:
            cell.font = HDR
            cell.fill = FILL
            cell.alignment = Alignment(horizontal="center")
        rows = (db.query(m.KassaEntry).filter(shart)
                .order_by(m.KassaEntry.entry_at, m.KassaEntry.id).all())
        kirim = chiqim = 0.0
        for e in rows:
            amt = float(e.amount)
            sfx = " $" if e.currency == "USD" else ""
            if e.direction == "Chiqim":
                ws.append([e.entry_at.isoformat(), e.who, e.note,
                           (amt if not sfx else f"{amt:g}{sfx}"), "", "", ""])
                if not sfx:
                    chiqim += amt
            else:
                ws.append([e.entry_at.isoformat(), "", "", "", e.who, e.note,
                           (amt if not sfx else f"{amt:g}{sfx}")])
                if not sfx:
                    kirim += amt
        ws.append([])
        ws.append(["", "", "JAMI CHIQIM:", chiqim, "", "JAMI KIRIM:", kirim])
        ws.append(["", "", "", "", "", "BALANS:", kirim - chiqim])
        for r in (ws.max_row - 1, ws.max_row):
            for cell in ws[r]:
                cell.font = Font(bold=True)
        for col in "DG":
            for cc in ws[col]:
                if isinstance(cc.value, (int, float)):
                    cc.number_format = "#,##0"
        for col, w in [("B", 22), ("C", 20), ("E", 18), ("F", 18)]:
            ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="kassa_{date.today()}.xlsx"'})
