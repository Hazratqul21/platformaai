#!/usr/bin/env python3
"""BILYARD KLUBI — Excel kassa daftarini tizimga kiritish.

Mijoz («THE BILLIARD») pulni Excel da yuritadi: har varaq — bir tur
harakat. Bu vosita o'sha varaqlarni o'qib platformaning KASSA
JURNALIGA (`kassa_entries`) yozadi — kirim/chiqim, sana, kim, izoh.

    .venv/bin/python tools/bilyard_excel.py "THE BILLIARD — август 2.xlsx" --tekshir
    .venv/bin/python tools/bilyard_excel.py "THE BILLIARD — август 2.xlsx"

    --tekshir   hech narsa yozmaydi, faqat nima kirishini ko'rsatadi
    --oy 2026-08  sanasi yo'q yozuvlar shu oyning 1-kuniga yoziladi
    --tozala    oldin kiritilganini o'chirib, qaytadan kiritadi

╔══════════════════════════════════════════════════════════════════╗
║  EXCEL FAYLGA YOZILMAYDI — `read_only=True` bilan ochiladi.      ║
╚══════════════════════════════════════════════════════════════════╝

VARAQLAR VA ULARNING MA'NOSI (fayl tekshirilib aniqlangan):

| Varaq       | Nima              | Yo'nalish |
|-------------|-------------------|-----------|
| ПРИХОД      | tushum            | Kirim     |
| Расходы     | kundalik xarajat  | Chiqim    |
| Запрлата    | oylik doimiy xarajat (ijara, svet, oylik) | Chiqim |
| Лист1       | bar uchun tovar   | Chiqim    |
| Лист2       | mijozning QO'LDA yozgan jamlari | O'TKAZILADI |
| 3-ЭТ 21 КВ  | boshqa loyihadan qolgan jamlovchi varaq | O'TKAZILADI |

Oxirgi ikkisi ATAYLAB olinmaydi: ular boshqa varaqlarning JAMI si.
Kiritilsa pul ikki marta sanaladi va kassa qoldig'i ikki barobar
noto'g'ri chiqadi.

TAKROR KIRITISHDAN HIMOYA: bazada shu vosita kiritgan yozuv bo'lsa,
vosita to'xtaydi. Qaytadan kiritish uchun `--tozala` beriladi — u
FAQAT shu vosita yozganini o'chiradi, qo'lda kiritilganiga tegmaydi.
"""
import argparse
import datetime
import re
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.db import SessionLocal                      # noqa: E402
from app import models as m                          # noqa: E402

KOK, SARIQ, QIZIL, TUGA = "\033[92m", "\033[93m", "\033[91m", "\033[0m"

BELGI = "Excel (bilyard)"      # `created_by` — qaysi yozuv qayerdan kelgani

# JAMLOVCHI QATORLAR — yozuv emas, yig'indi. Kiritilsa pul ikki marta
# sanaladi va qoldiq ikki barobar noto'g'ri chiqadi.
#
# Ikki xil tekshiruv, chunki ikki xil ko'rinishda uchraydi:
#   · ichida «Итого» bo'lgan matn      — «Итого прямые затраты»
#   · nomi AYNAN tur so'zi bo'lgan qator — «зарплата» deb yozilgan
#     qator `Запрлата` varag'ining JAMI si (104 340 000) edi va u
#     xarajatni ikki barobar ko'rsatardi. Faylda ushlandi.
JAMI_ICHIDA = ("итого", "всего", "итог", "jami", "жами")
JAMI_AYNAN = {"приход", "расход", "расходы", "зарплата", "зарплат",
              "приходы", "остаток", "итого"}

# varaq: (nom_ustuni, tur_ustuni, summa_ustuni, usd_ustuni, sana_ustuni, yonalish)
# Ustunlar 0 dan sanaladi, birinchi qator — sarlavha.
VARAQLAR = {
    "ПРИХОД":   (1, None, 2, 3, 4, "Kirim"),
    "Запрлата": (1, 2,    5, 6, 7, "Chiqim"),
    "Расходы":  (1, 2,    5, 6, 7, "Chiqim"),
    "Лист1":    (0, 1,    4, 5, 6, "Chiqim"),
}
OTKAZILADI = ("Лист2", "3-ЭТ 21 КВ")


def _sana(qiymat, zaxira: datetime.date) -> tuple[datetime.date, bool]:
    """Sanani o'qiydi. Ikki shakl uchraydi: haqiqiy sana va «01,08,2026».

    Qaytaradi: (sana, haqiqiymi). Sana bo'lmasa `zaxira` beriladi va
    `False` — bu yozuv «oyning boshiga qo'yildi» deb sanaladi.
    """
    if isinstance(qiymat, datetime.datetime):
        return qiymat.date(), True
    if isinstance(qiymat, datetime.date):
        return qiymat, True
    if isinstance(qiymat, str):
        mos = re.match(r"\s*(\d{1,2})[,.\-/](\d{1,2})[,.\-/](\d{4})", qiymat)
        if mos:
            kun, oy, yil = (int(x) for x in mos.groups())
            try:
                return datetime.date(yil, oy, kun), True
            except ValueError:
                pass
    return zaxira, False


def _summa(qiymat) -> Decimal:
    if qiymat in (None, ""):
        return Decimal("0")
    if isinstance(qiymat, str):
        qiymat = qiymat.replace(" ", "").replace("\xa0", "").replace(",", ".")
    try:
        return Decimal(str(qiymat)).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _matn(qiymat, uzunlik: int) -> str:
    return ("" if qiymat is None else str(qiymat)).strip()[:uzunlik]


def _jamimi(nom: str) -> bool:
    past = nom.strip().lower()
    return past in JAMI_AYNAN or any(soz in past for soz in JAMI_ICHIDA)


def kirit(fayl: str, tekshir: bool = False, oy: str = "",
          tozala: bool = False) -> dict:
    import openpyxl

    yol = Path(fayl).expanduser().resolve()
    if not yol.exists():
        sys.exit(f"❌ Fayl topilmadi: {yol}")

    if oy:
        try:
            zaxira_sana = datetime.datetime.strptime(oy + "-01", "%Y-%m-%d").date()
        except ValueError:
            sys.exit("❌ --oy shakli: 2026-08")
    else:
        zaxira_sana = None            # pastda faylning o'zidan aniqlanadi

    kitob = openpyxl.load_workbook(yol, data_only=True, read_only=True)
    db = SessionLocal()
    hisob = {"kiritildi": 0, "otkazildi": 0, "sanasiz": 0}
    jami = {"Kirim": Decimal("0"), "Chiqim": Decimal("0")}
    varaq_hisobi = {}

    try:
        # 1-yurish: sanasi bor yozuvlardan oyni aniqlaymiz (agar
        # `--oy` berilmagan bo'lsa). Sanasiz yozuvlar (ijara, oylik)
        # o'sha oyning 1-kuniga qo'yiladi.
        if zaxira_sana is None:
            sanalar = []
            for varaq, (n_u, _t, s_u, _u, d_u, _y) in VARAQLAR.items():
                if varaq not in kitob.sheetnames:
                    continue
                for qator in kitob[varaq].iter_rows(min_row=2, values_only=True):
                    if len(qator) <= max(n_u, s_u, d_u):
                        continue
                    if _summa(qator[s_u]) <= 0:
                        continue
                    d, haqiqiy = _sana(qator[d_u], datetime.date.today())
                    if haqiqiy:
                        sanalar.append(d)
            if sanalar:
                eng_kop = max(set((d.year, d.month) for d in sanalar),
                              key=lambda ym: sum(1 for d in sanalar
                                                 if (d.year, d.month) == ym))
                zaxira_sana = datetime.date(eng_kop[0], eng_kop[1], 1)
            else:
                zaxira_sana = datetime.date.today().replace(day=1)
            print(f"Sanasiz yozuvlar uchun: {zaxira_sana} (fayldan aniqlandi)")

        # TAKROR KIRITISHDAN HIMOYA — fayl darajasida, qator darajasida
        # EMAS.
        #
        # Avval qator darajasida edi: bir xil (sana+kim+summa) yozuv
        # ikkinchi marta uchrasa tashlab yuborilardi. Bu XATO edi —
        # faylda «Алиш ишчи · 50 000 · 07.08» yozuvi ikki marta bor va
        # ikkalasi ham HAQIQIY (ikki kishiga berilgan premiya). Ya'ni
        # himoya haqiqiy pulni yeb qo'yardi.
        #
        # Endi: shu vosita kiritgan yozuv bazada bormi — bo'lsa
        # to'xtaydi. Qayta kiritish kerak bo'lsa `--tozala` beriladi,
        # u FAQAT shu vosita yozganini o'chiradi (qo'lda kiritilgan
        # yozuvga tegmaydi).
        eski = db.query(m.KassaEntry).filter(
            m.KassaEntry.created_by == BELGI).count()
        if eski and tozala:
            db.query(m.KassaEntry).filter(
                m.KassaEntry.created_by == BELGI).delete()
            db.flush()
            print(f"{SARIQ}○ Oldingi Excel yozuvlari o'chirildi: {eski} ta{TUGA}")
        elif eski:
            db.rollback()
            sys.exit(
                f"❌ Bu bazada Excel dan kiritilgan {eski} ta yozuv bor.\n"
                f"   Qayta kiritish uchun: --tozala (faqat Excel "
                f"yozuvlarini o'chiradi, qo'lda kiritilganiga tegmaydi)")

        for varaq in kitob.sheetnames:
            if varaq.strip() in OTKAZILADI or varaq not in VARAQLAR:
                if varaq.strip() in OTKAZILADI:
                    print(f"  ○ «{varaq.strip()}» o'tkazildi — jamlovchi varaq")
                continue

            n_u, t_u, s_u, u_u, d_u, yonalish = VARAQLAR[varaq]
            soni = 0
            varaq_jami = Decimal("0")

            for qator in kitob[varaq].iter_rows(min_row=2, values_only=True):
                if len(qator) <= max(n_u, s_u):
                    continue
                nom = _matn(qator[n_u], 120)
                summa = _summa(qator[s_u])
                usd = _summa(qator[u_u]) if u_u is not None and len(qator) > u_u \
                    else Decimal("0")
                if not nom or (summa <= 0 and usd <= 0):
                    continue
                if _jamimi(nom):
                    hisob["otkazildi"] += 1
                    continue

                # Dollarda yozilgan yozuv ham bor — valyutasi saqlanadi,
                # so'mga aylantirilmaydi (kurs noma'lum, taxmin qilinmaydi).
                if summa <= 0 and usd > 0:
                    summa, valyuta = usd, "USD"
                else:
                    valyuta = "so'm"

                sana, haqiqiy = _sana(qator[d_u] if len(qator) > d_u else None,
                                      zaxira_sana)
                if not haqiqiy:
                    hisob["sanasiz"] += 1

                tur = _matn(qator[t_u], 60) if t_u is not None and len(qator) > t_u \
                    else ""
                izoh = tur or varaq.strip()
                if not haqiqiy:
                    izoh = f"{izoh} · oylik (sanasiz)"

                db.add(m.KassaEntry(
                    firm="THE BILLIARD", direction=yonalish, who=nom,
                    note=izoh[:200], amount=summa, currency=valyuta,
                    entry_at=sana, created_by=BELGI))
                soni += 1
                hisob["kiritildi"] += 1
                if valyuta == "so'm":
                    jami[yonalish] += summa
                    varaq_jami += summa

            varaq_hisobi[varaq.strip()] = (soni, varaq_jami)
            print(f"  {varaq.strip():<12} {soni:>4} ta · {varaq_jami:>16,.0f} so'm"
                  f"  ({yonalish})")

        if tekshir:
            db.rollback()
            print(f"\n{SARIQ}🔍 TEKSHIRUV REJIMI — hech narsa yozilmadi{TUGA}")
        else:
            db.add(m.AuditLog(
                who=BELGI, action="Excel dan kassa yozuvlari kiritildi",
                detail=f"{yol.name}: " + ", ".join(
                    f"{k} {v[0]}" for k, v in varaq_hisobi.items())))
            db.commit()
    finally:
        db.close()

    hisob["kirim_jami"] = jami["Kirim"]
    hisob["chiqim_jami"] = jami["Chiqim"]
    return hisob


def main():
    p = argparse.ArgumentParser(description="Bilyard Excel kassa daftari")
    p.add_argument("fayl", help="Excel fayl (.xlsx)")
    p.add_argument("--tekshir", action="store_true",
                   help="Hech narsa yozmaydi, faqat ko'rsatadi")
    p.add_argument("--oy", default="",
                   help="Sanasiz yozuvlar uchun oy: 2026-08")
    p.add_argument("--tozala", action="store_true",
                   help="Oldin shu vosita kiritgan yozuvlarni o'chirib, "
                        "qaytadan kiritadi (qo'lda kiritilganiga tegmaydi)")
    a = p.parse_args()

    print(f"\nManba (faqat o'qish): {a.fayl}\n")
    h = kirit(a.fayl, a.tekshir, a.oy, a.tozala)

    print(f"\n{'='*54}")
    print(f"  kiritildi        {h['kiritildi']:>6}")
    print(f"  jamlovchi qator  {h['otkazildi']:>6}  (yig'indi — olinmadi)")
    print(f"  sanasiz          {h['sanasiz']:>6}  (oy boshiga qo'yildi)")
    print(f"{'-'*54}")
    print(f"  KIRIM   {h['kirim_jami']:>18,.0f} so'm")
    print(f"  CHIQIM  {h['chiqim_jami']:>18,.0f} so'm")
    farq = h["kirim_jami"] - h["chiqim_jami"]
    rang = KOK if farq >= 0 else QIZIL
    print(f"  {rang}QOLDIQ  {farq:>18,.0f} so'm{TUGA}")
    print(f"{'='*54}")
    if not a.tekshir:
        print(f"{KOK}✅ Kassa jurnaliga yozildi{TUGA}")


if __name__ == "__main__":
    main()
